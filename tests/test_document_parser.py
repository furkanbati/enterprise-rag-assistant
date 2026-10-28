
from types import SimpleNamespace

import docx
import pytest
from docling.datamodel.base_models import InputFormat
from docling_core.types.doc import DocItemLabel

from document_parser import DocumentParser


class FakeDocument:
    def __init__(self, items):
        self.items = items

    def iterate_items(self):
        for item in self.items:
            yield item, None


class FakeConverter:
    def __init__(self, items):
        self.document = FakeDocument(items)
        self.received_path = None

    def convert(self, pdf_path):
        self.received_path = pdf_path
        return SimpleNamespace(document=self.document)


def make_item(
    label,
    text,
    level=1,
    page=1,
    with_provenance=True,
):
    provenance = (
        [SimpleNamespace(page_no=page)]
        if with_provenance
        else []
    )

    return SimpleNamespace(
        label=label,
        text=text,
        level=level,
        prov=provenance,
    )


def make_parser(items, ocr_items=()):
    parser = DocumentParser.__new__(DocumentParser)
    parser.converter = FakeConverter(items)
    parser.ocr_converter = FakeConverter(list(ocr_items))
    return parser


def test_parse_retries_a_pdf_without_text_with_ocr():
    parser = make_parser(
        [],
        ocr_items=[
            make_item(DocItemLabel.SECTION_HEADER, "Scanned policy", page=1),
            make_item(DocItemLabel.TEXT, "Text read by OCR.", page=1),
        ],
    )

    sections = parser.parse("scan.pdf")

    assert [(s.heading, s.page, s.text) for s in sections] == [
        ("Scanned policy", 1, "Text read by OCR."),
    ]
    assert parser.ocr_converter.received_path == "scan.pdf"


def test_parse_skips_ocr_when_the_pdf_has_text():
    parser = make_parser(
        [make_item(DocItemLabel.TEXT, "Digital text.")],
        ocr_items=[make_item(DocItemLabel.TEXT, "OCR text.")],
    )

    sections = parser.parse("guide.pdf")

    assert [s.text for s in sections] == ["Digital text."]
    assert parser.ocr_converter.received_path is None


def test_parse_does_not_run_ocr_for_other_formats():
    parser = make_parser(
        [],
        ocr_items=[make_item(DocItemLabel.TEXT, "OCR text.")],
    )

    assert parser.parse("empty.docx") == []
    assert parser.ocr_converter.received_path is None


def test_parse_returns_nothing_for_a_pdf_without_text_when_ocr_is_off():
    parser = make_parser([])
    parser.ocr_converter = None

    assert parser.parse("scan.pdf") == []


@pytest.mark.parametrize("engine", ["rapidocr", "easyocr"])
def test_ocr_converter_uses_the_configured_engine(engine):
    parser = DocumentParser(ocr_engine=engine)

    options = parser.ocr_converter.format_to_options[InputFormat.PDF].pipeline_options

    assert options.do_ocr
    assert options.ocr_options.kind == engine
    assert options.ocr_options.force_full_page_ocr


def test_ocr_can_be_turned_off():
    assert DocumentParser(ocr_engine="off").ocr_converter is None


def test_unknown_ocr_engine_is_rejected():
    with pytest.raises(ValueError, match="Unknown OCR engine 'tesseract'"):
        DocumentParser(ocr_engine="tesseract")


def test_parse_groups_text_under_section_headings():
    parser = make_parser(
        [
            make_item(
                DocItemLabel.SECTION_HEADER,
                "Introduction",
                level=1,
                page=2,
            ),
            make_item(
                DocItemLabel.TEXT,
                "First paragraph.",
                page=2,
            ),
            make_item(
                DocItemLabel.SECTION_HEADER,
                "Methods",
                level=2,
                page=3,
            ),
            make_item(
                DocItemLabel.TEXT,
                "Method description.",
                page=3,
            ),
        ]
    )

    sections = parser.parse("guide.pdf")

    assert len(sections) == 2
    assert sections[0].heading == "Introduction"
    assert sections[0].level == 1
    assert sections[0].page == 2
    assert sections[0].text == "First paragraph."

    assert sections[1].heading == "Methods"
    assert sections[1].level == 2
    assert sections[1].page == 3
    assert sections[1].text == "Method description."
    assert parser.converter.received_path == "guide.pdf"


def test_parse_formats_list_items_as_bullets():
    parser = make_parser(
        [
            make_item(
                DocItemLabel.SECTION_HEADER,
                "Features",
            ),
            make_item(
                DocItemLabel.LIST_ITEM,
                "Search documents",
            ),
            make_item(
                DocItemLabel.LIST_ITEM,
                "Generate answers",
            ),
        ]
    )

    sections = parser.parse("guide.pdf")

    assert len(sections) == 1
    assert sections[0].text == (
        "- Search documents\n- Generate answers"
    )


def test_parse_uses_default_page_when_heading_has_no_provenance():
    parser = make_parser(
        [
            make_item(
                DocItemLabel.SECTION_HEADER,
                "Overview",
                with_provenance=False,
            ),
            make_item(
                DocItemLabel.TEXT,
                "Overview text.",
            ),
        ]
    )

    sections = parser.parse("guide.pdf")

    assert len(sections) == 1
    assert sections[0].page == 1


def test_parse_returns_empty_list_when_document_has_no_text():
    parser = make_parser(
        [
            make_item(
                DocItemLabel.SECTION_HEADER,
                "Empty section",
            ),
        ]
    )

    sections = parser.parse("empty.pdf")

    assert sections == []


def test_parse_treats_titles_as_headings_and_paragraphs_as_text():
    parser = make_parser(
        [
            make_item(DocItemLabel.TITLE, "Handbook", level=None, with_provenance=False),
            make_item(DocItemLabel.PARAGRAPH, "Intro paragraph.", with_provenance=False),
        ]
    )

    sections = parser.parse("handbook.docx")

    assert [(s.heading, s.level, s.page, s.text) for s in sections] == [
        ("Handbook", 1, 1, "Intro paragraph."),
    ]


def test_parse_reads_plain_text_without_docling(tmp_path):
    parser = make_parser([])
    path = tmp_path / "notes.txt"
    path.write_text("First line.\nSecond line.\n", encoding="utf-8")

    sections = parser.parse(str(path))

    assert [(s.heading, s.page, s.text) for s in sections] == [
        ("", 1, "First line.\nSecond line."),
    ]
    assert parser.converter.received_path is None


def test_parse_returns_nothing_for_an_empty_text_file(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("  \n", encoding="utf-8")

    assert make_parser([]).parse(str(path)) == []


@pytest.mark.parametrize("name", ["policy.md", "policy.html", "policy.docx"])
def test_parse_real_documents_in_each_supported_format(tmp_path, name):
    path = tmp_path / name

    if name.endswith(".md"):
        path.write_text(
            "# Handbook\n\nIntro paragraph.\n\n## Leave policy\n\nEmployees get 20 days.\n",
            encoding="utf-8",
        )
    elif name.endswith(".html"):
        path.write_text(
            "<html><body><h1>Handbook</h1><p>Intro paragraph.</p>"
            "<h2>Leave policy</h2><p>Employees get 20 days.</p></body></html>",
            encoding="utf-8",
        )
    else:
        document = docx.Document()
        document.add_heading("Handbook", 0)
        document.add_paragraph("Intro paragraph.")
        document.add_heading("Leave policy", 1)
        document.add_paragraph("Employees get 20 days.")
        document.save(path)

    sections = DocumentParser().parse(str(path))

    assert [(s.heading, s.text) for s in sections] == [
        ("Handbook", "Intro paragraph."),
        ("Leave policy", "Employees get 20 days."),
    ]
